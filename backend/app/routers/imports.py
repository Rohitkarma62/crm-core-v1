import csv, io, re, json
from decimal import Decimal, InvalidOperation
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy import and_, func, or_, select
from sqlalchemy.orm import Session
from ..database import get_db
from ..dependencies import get_current_user, require_permission
from ..models.core import User
from ..models.crm import Customer, Lead, LeadSource, LeadStatus, Payment, Sale, ImportJob
from ..schemas.imports import ImportOptions, ImportResult

router = APIRouter(prefix="/api/v1/imports", tags=["Import / Export"])
MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10000
FIELD_ALIASES = {
    "name": {"name", "full name", "client name", "customer name", "lead name"},
    "phone": {"phone", "mobile", "mobile number", "phone number", "contact"},
    "email": {"email", "e-mail", "email address", "e mail"},
    "company": {"company", "business", "company name"},
    "source": {"source", "lead source"},
    "status": {"status", "lead status"},
    "priority": {"priority"},
    "interested_service": {"service", "interested service", "product", "interested product"},
    "estimated_value": {"value", "estimated value", "amount", "deal value"},
    "notes": {"notes", "note", "remarks", "description"},
    "assigned_to": {"assigned to", "assignee", "staff", "salesperson"},
}
ALLOWED_FIELDS = set(FIELD_ALIASES)

def normalize_header(value):
    return re.sub(r"\s+", " ", str(value or "").strip().lower().replace("_", " ").replace("-", " "))

def normalize_phone(value):
    return re.sub(r"\D", "", str(value or ""))

def parse_value(value):
    if value is None: return None
    text = str(value).strip().replace(",", "").replace("₹", "").strip()
    if not text: return None
    try: return Decimal(text)
    except InvalidOperation: return None

def suggest_mapping(headers):
    mapping = {}
    for header in headers:
        norm = normalize_header(header)
        for field, aliases in FIELD_ALIASES.items():
            if norm in aliases:
                mapping[header] = field
                break
    return mapping

def read_rows(filename, content):
    if len(content) > MAX_BYTES: raise HTTPException(413, "File exceeds 10 MB limit")
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "csv":
        try: text = content.decode("utf-8-sig")
        except UnicodeDecodeError: text = content.decode("latin-1")
        reader = csv.DictReader(io.StringIO(text))
        rows = list(reader)
        headers = reader.fieldnames or []
    elif ext in {"xlsx", "xls"}:
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise HTTPException(500, "Excel import dependency is not installed")
        if ext == "xls": raise HTTPException(400, "Legacy .xls is not supported in V1; save it as .xlsx")
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        values = list(ws.iter_rows(values_only=True))
        if not values: return [], []
        headers = [str(x or "").strip() for x in values[0]]
        rows = [dict(zip(headers, row)) for row in values[1:] if any(x is not None and str(x).strip() for x in row)]
    else:
        raise HTTPException(400, "Only CSV and XLSX files are supported")
    if len(rows) > MAX_ROWS: raise HTTPException(400, f"Maximum {MAX_ROWS} rows per import")
    return headers, rows

@router.post("/leads/preview")
async def preview_leads(file: UploadFile = File(...), db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    require_permission(user, "leads.import", db)
    content = await file.read()
    headers, rows = read_rows(file.filename or "upload", content)
    mapping = suggest_mapping(headers)
    missing = [field for field in ("name",) if field not in mapping.values()]
    job = ImportJob(
        business_id=user.business_id, user_id=user.id, filename=file.filename or "upload", entity="leads",
        total_rows=len(rows), staged_rows=json.dumps(rows, default=str), staged_headers=json.dumps(headers), status="staged"
    )
    db.add(job); db.commit(); db.refresh(job)
    return {"job_id": job.id, "headers": headers, "suggested_mapping": mapping, "required_missing": missing, "total_rows": len(rows), "sample_rows": rows[:10]}

@router.post("/leads/import", response_model=ImportResult)
def import_leads(payload: ImportOptions, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    require_permission(user, "leads.import", db)
    if not payload.job_id:
        raise HTTPException(400, "job_id is required. Upload and preview the file first.")
    if payload.duplicate_action not in {"skip", "update", "create"}:
        raise HTTPException(400, "duplicate_action must be skip, update, or create")
    job = db.scalar(select(ImportJob).where(ImportJob.id == payload.job_id, ImportJob.business_id == user.business_id))
    if not job: raise HTTPException(404, "Import job not found")
    if job.status != "staged": raise HTTPException(409, "Import job has already been processed")
    if not payload.mapping: raise HTTPException(400, "Column mapping is required")
    try: rows = json.loads(job.staged_rows or "[]")
    except json.JSONDecodeError: raise HTTPException(500, "Staged import data is invalid")
    if len(rows) > MAX_ROWS: raise HTTPException(400, f"Maximum {MAX_ROWS} rows per import")
    if "name" not in set(payload.mapping.values()): raise HTTPException(400, "Name mapping is required")

    sources = {s.name.strip().lower(): s for s in db.scalars(select(LeadSource).where(LeadSource.business_id == user.business_id)).all()}
    statuses = {s.name.strip().lower(): s for s in db.scalars(select(LeadStatus).where(LeadStatus.business_id == user.business_id)).all()}
    users = {u.email.strip().lower(): u for u in db.scalars(select(User).where(User.business_id == user.business_id)).all() if u.email}
    imported = skipped = errors = 0; error_rows = []
    for idx, raw in enumerate(rows, start=2):
        data = {field: raw.get(col) for col, field in payload.mapping.items() if field in ALLOWED_FIELDS}
        name = str(data.get("name") or "").strip(); phone = normalize_phone(data.get("phone"))
        if not name: errors += 1; error_rows.append({"row": idx, "error": "Name is required"}); continue
        if not phone: errors += 1; error_rows.append({"row": idx, "error": "Phone is required"}); continue
        if len(phone) < 7: errors += 1; error_rows.append({"row": idx, "error": "Invalid phone"}); continue
        email = str(data.get("email") or "").strip() or None
        duplicate = db.scalar(select(Lead).where(Lead.business_id == user.business_id, or_(Lead.phone == phone, Lead.email == email) if email else Lead.phone == phone))
        if duplicate:
            if payload.duplicate_action == "skip": skipped += 1; continue
            lead = duplicate if payload.duplicate_action == "update" else None
        else: lead = None
        source_name = str(data.get("source") or "").strip(); source_id = None
        if source_name:
            source = sources.get(source_name.lower())
            if not source: source = LeadSource(business_id=user.business_id, name=source_name); db.add(source); db.flush(); sources[source_name.lower()] = source
            source_id = source.id
        status_name = str(data.get("status") or "").strip(); status_id = None
        if status_name:
            status = statuses.get(status_name.lower())
            if not status: status = LeadStatus(business_id=user.business_id, name=status_name, sort_order=0); db.add(status); db.flush(); statuses[status_name.lower()] = status
            status_id = status.id
        priority = str(data.get("priority") or "medium").strip().lower()
        if priority not in {"low", "medium", "high"}: priority = "medium"
        assigned = None; assignee = str(data.get("assigned_to") or "").strip().lower()
        if assignee and assignee in users: assigned = users[assignee].id
        estimated = parse_value(data.get("estimated_value"))
        if estimated is None and str(data.get("estimated_value") or "").strip():
            errors += 1; error_rows.append({"row": idx, "error": "Invalid estimated value"}); continue
        values = {"business_id": user.business_id, "name": name, "phone": phone, "email": email, "company": str(data.get("company") or "").strip() or None, "source_id": source_id, "status_id": status_id, "priority": priority, "interested_service": str(data.get("interested_service") or "").strip() or None, "estimated_value": estimated, "notes": str(data.get("notes") or "").strip() or None, "assigned_to": assigned}
        if lead:
            for key, value in values.items():
                if key != "business_id": setattr(lead, key, value)
        else: db.add(Lead(**values))
        imported += 1
    job.imported=imported; job.skipped=skipped; job.errors=errors; job.error_rows=json.dumps(error_rows[:500], default=str); job.status="completed"
    db.commit()
    return ImportResult(total_rows=len(rows), imported=imported, skipped=skipped, errors=errors, error_rows=error_rows[:500])

def _safe_cell(value):
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value

def _csv_response(rows, filename):
    buf = io.StringIO(); writer = csv.writer(buf)
    if rows:
        writer.writerow(rows[0].keys())
        for row in rows: writer.writerow([_safe_cell(row.get(k)) for k in rows[0].keys()])
    data = buf.getvalue().encode("utf-8-sig")
    return StreamingResponse(io.BytesIO(data), media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="{filename}"'})

def _xlsx_response(rows, filename):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    if rows:
        keys = list(rows[0].keys()); ws.append(keys)
        for row in rows: ws.append([_safe_cell(row.get(k)) for k in keys])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename}"'})

def _export(rows, fmt, filename):
    fmt = (fmt or "csv").lower()
    if fmt not in {"csv", "xlsx"}:
        raise HTTPException(400, "fmt must be csv or xlsx")
    if fmt == "xlsx": return _xlsx_response(rows, filename + ".xlsx")
    return _csv_response(rows, filename + ".csv")

@router.get("/leads/export")
def export_leads(fmt: str = "csv", search: str | None = None, status: str | None = None, priority: str | None = None, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    require_permission(user, "leads.export", db)
    conditions = [Lead.business_id == user.business_id]
    if search:
        term=f"%{search.strip()}%"; conditions.append(or_(Lead.name.ilike(term), Lead.phone.ilike(term), Lead.email.ilike(term), Lead.company.ilike(term)))
    if priority: conditions.append(Lead.priority == priority)
    stmt = select(Lead).where(*conditions).order_by(Lead.created_at.desc())
    leads = db.scalars(stmt).all()
    status_map = {s.id:s.name for s in db.scalars(select(LeadStatus).where(LeadStatus.business_id==user.business_id)).all()}
    source_map = {s.id:s.name for s in db.scalars(select(LeadSource).where(LeadSource.business_id==user.business_id)).all()}
    if status: leads=[x for x in leads if status_map.get(x.status_id, "").lower()==status.lower()]
    rows=[{"id":x.id,"name":x.name,"phone":x.phone,"email":x.email or "","company":x.company or "","source":source_map.get(x.source_id,""),"status":status_map.get(x.status_id,""),"priority":x.priority,"interested_service":x.interested_service or "","estimated_value":str(x.estimated_value or ""),"notes":x.notes or "","created_at":x.created_at.isoformat()} for x in leads]
    return _export(rows, fmt, "crm-leads")

@router.get("/customers/export")
def export_customers(fmt: str = "csv", search: str | None = None, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    require_permission(user, "customers.export", db)
    conditions=[Customer.business_id==user.business_id]
    if search:
        term=f"%{search.strip()}%"; conditions.append(or_(Customer.name.ilike(term), Customer.phone.ilike(term), Customer.email.ilike(term), Customer.company.ilike(term)))
    items=db.scalars(select(Customer).where(*conditions).order_by(Customer.created_at.desc())).all()
    rows=[{"id":x.id,"name":x.name,"phone":x.phone,"email":x.email or "","company":x.company or "","address":x.address or "","created_at":x.created_at.isoformat()} for x in items]
    return _export(rows, fmt, "crm-customers")

@router.get("/sales/export")
def export_sales(fmt: str = "csv", db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    require_permission(user, "sales.export", db)
    sales=db.scalars(select(Sale).where(Sale.business_id==user.business_id).order_by(Sale.sale_date.desc())).all()
    customers={c.id:c.name for c in db.scalars(select(Customer).where(Customer.business_id==user.business_id)).all()}
    rows=[{"id":x.id,"customer":customers.get(x.customer_id,""),"amount":str(x.amount),"status":x.status,"sale_date":x.sale_date.isoformat(),"notes":x.notes or ""} for x in sales]
    return _export(rows, fmt, "crm-sales")

@router.get("/payments/export")
def export_payments(fmt: str = "csv", db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    require_permission(user, "payments.export", db)
    payments=db.scalars(select(Payment).where(Payment.business_id==user.business_id).order_by(Payment.payment_date.desc())).all()
    customers={c.id:c.name for c in db.scalars(select(Customer).where(Customer.business_id==user.business_id)).all()}
    rows=[{"id":x.id,"customer":customers.get(x.customer_id,""),"sale_id":x.sale_id,"amount":str(x.amount),"payment_method":x.payment_method or "","payment_date":x.payment_date.isoformat(),"status":x.status,"reference":x.reference or ""} for x in payments]
    return _export(rows, fmt, "crm-payments")

@router.get("/history")
def import_history(limit: int = 50, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    require_permission(user, "leads.import", db)
    limit = max(1, min(limit, 100))
    jobs=db.scalars(select(ImportJob).where(ImportJob.business_id==user.business_id).order_by(ImportJob.created_at.desc()).limit(limit)).all()
    return [{"id":j.id,"filename":j.filename,"entity":j.entity,"total_rows":j.total_rows,"imported":j.imported,"skipped":j.skipped,"errors":j.errors,"status":j.status,"created_at":j.created_at.isoformat()} for j in jobs]

@router.delete("/history/{job_id}")
def delete_import_job(job_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    require_permission(user, "leads.import", db)
    job=db.scalar(select(ImportJob).where(ImportJob.id==job_id, ImportJob.business_id==user.business_id))
    if not job: raise HTTPException(404,"Import job not found")
    db.delete(job); db.commit(); return {"ok": True}

@router.get("/history/{job_id}/errors")
def import_errors(job_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    require_permission(user, "leads.import", db)
    job=db.scalar(select(ImportJob).where(ImportJob.id==job_id, ImportJob.business_id==user.business_id))
    if not job: raise HTTPException(404,"Import job not found")
    try: errors=json.loads(job.error_rows or "[]")
    except json.JSONDecodeError: errors=[]
    rows=[{"row":e.get("row"),"error":e.get("error")} for e in errors]
    return _csv_response(rows, f"import-{job_id}-errors.csv")
